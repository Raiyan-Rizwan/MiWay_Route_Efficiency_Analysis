from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "Outputs"
REPORT_DIR = PROJECT_ROOT / "Reports"
REPORT_DIR.mkdir(exist_ok=True)

DOC_PATH = REPORT_DIR / "MiWay_Analysis_Pipeline_Document.docx"
ANALYSIS_DATE = "Tuesday, May 5, 2026"
FEED_PERIOD = "May 1, 2026 to June 28, 2026"


def set_run_font(run, size=None, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9E2EC", size="4"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    total_dxa = int(sum(widths) * 1440)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_dxa))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Calibri")
    normal._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Calibri")
    title._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("0B2545")

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 18, 10),
        ("Heading 2", 13, "2E74B5", 14, 7),
        ("Heading 3", 12, "1F4D78", 10, 5),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Calibri")
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("MiWay Analysis Pipeline")
    set_run_font(run, size=9, color="666666")


def add_paragraph(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest)
    else:
        run = p.add_run(text)
        set_run_font(run)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        run = p.add_run(item)
        set_run_font(run)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        run = p.add_run(item)
        set_run_font(run)


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    set_cell_border(cell, color="B7C9DB", size="6")
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(title)
    set_run_font(run, size=11, color="0B2545", bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    run = p2.add_run(body)
    set_run_font(run, size=10.5)
    doc.add_paragraph()


def add_table(doc, rows, headers, widths, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_width(table, widths)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, "E8EEF5")
        set_cell_margins(cell)
        set_cell_border(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        set_run_font(run, size=font_size, color="0B2545", bold=True)

    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cell = cells[idx]
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if idx > 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(str(value))
            set_run_font(run, size=font_size)
    doc.add_paragraph()
    return table


def add_df_table(doc, df, columns, headers, widths, font_size=8):
    rows = []
    for _, row in df.iterrows():
        out = []
        for col in columns:
            value = row[col]
            if pd.isna(value):
                out.append("n/a")
            elif isinstance(value, float):
                out.append(f"{value:.2f}")
            else:
                out.append(str(value))
        rows.append(out)
    return add_table(doc, rows, headers, widths, font_size=font_size)


def load_outputs():
    ranking = pd.read_csv(OUTPUT_DIR / "final_route_efficiency_ranking_2026-05-05.csv")
    priority = pd.read_csv(OUTPUT_DIR / "advocacy_priority_routes_2026-05-05.csv")
    components = pd.read_csv(OUTPUT_DIR / "route_efficiency_score_components_2026-05-05.csv")
    geometry = pd.read_csv(OUTPUT_DIR / "route_geometry_stop_metrics_2026-05-05.csv")
    speed = pd.read_csv(OUTPUT_DIR / "route_scheduled_speed_metrics_2026-05-05.csv")
    frequency = pd.read_csv(OUTPUT_DIR / "route_frequency_service_metrics_2026-05-05.csv")
    return ranking, priority, components, geometry, speed, frequency


def add_image_if_exists(doc, path, caption):
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(path), width=Inches(6.3))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap_run = cap.add_run(caption)
    set_run_font(cap_run, size=9, color="555555", italic=True)


def build_doc():
    ranking, priority, components, geometry, speed, frequency = load_outputs()

    doc = Document()
    style_document(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("MiWay Route Efficiency and Advocacy Priority Analysis Pipeline")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("End-to-end analytical workflow from GTFS feed to route scores, priority routes, and ArcGIS map outputs")
    set_run_font(run, size=12, color="555555")

    meta_rows = [
        ["Project", "MiWay Route Efficiency"],
        ["Prepared for", "UTM Transit Housing Advocacy"],
        ["Input data", "MiWay GTFS scheduled transit feed"],
        ["Feed period", FEED_PERIOD],
        ["Representative service day", ANALYSIS_DATE],
        ["Primary outputs", "Final route ranking CSV, advocacy priority CSV, ArcGIS route-line GeoJSONs, map layouts"],
    ]
    add_table(doc, meta_rows, ["Field", "Value"], [1.7, 4.8], font_size=9)

    add_callout(
        doc,
        "Purpose of the pipeline",
        "This pipeline turns raw scheduled transit data into a transparent, repeatable route-efficiency scoring system. "
        "It separates operational efficiency from advocacy priority so the project can identify strong corridors, weaker corridors, "
        "and routes where improvement may matter because the route has meaningful scheduled service volume.",
    )

    doc.add_heading("1. Pipeline Overview", level=1)
    add_paragraph(
        doc,
        "The analysis is organized as five notebooks plus one ArcGIS handoff script. Each notebook produces route-level outputs that become inputs to the next stage.",
    )
    overview_rows = [
        ["1", "GTFS exploration", "Load feed, inspect tables, validate dates, select active service for the representative weekday."],
        ["2", "Route geometry", "Calculate shape lengths, directness ratios, unique stops, average stops per trip, and stops per kilometre."],
        ["3", "Scheduled speed", "Parse GTFS times, compute trip durations, attach shape distances, and summarize scheduled route speeds."],
        ["4", "Frequency and availability", "Calculate trips, headways, service spans, period-level service, directional balance, and all-day flags."],
        ["5", "Final scoring", "Normalize metrics into 0-100 scores, calculate final efficiency scores, rank routes, and create advocacy priorities."],
        ["6", "ArcGIS handoff", "Join final scores to route geometry and export ArcGIS-ready route-line layers for mapping."],
    ]
    add_table(doc, overview_rows, ["Stage", "Name", "What it does"], [0.65, 1.55, 4.3], font_size=8.5)

    doc.add_heading("2. Source Data and Core Tables", level=1)
    add_paragraph(
        doc,
        "The project starts from a GTFS feed stored at Data/google_transit.zip and extracted into Data/google_transit. GTFS is useful because it provides a standardized scheduled-service model: routes, trips, stops, stop times, service calendars, and route shapes.",
    )
    gtfs_rows = [
        ["routes.txt", "Route IDs, short names, long names, colors", "Used to label route-level outputs and map legends."],
        ["trips.txt", "Trip IDs, route IDs, service IDs, direction IDs, shape IDs", "Connects scheduled trips to routes and route shapes."],
        ["stop_times.txt", "Stop sequence, arrival/departure times, stop IDs", "Used for trip duration, stop counts, and departure-based headways."],
        ["stops.txt", "Stop coordinates and stop names", "Used for stop-level cleaning and optional ArcGIS stop layers."],
        ["shapes.txt", "Ordered latitude/longitude points for each path", "Used to calculate path length and create route lines."],
        ["calendar_dates.txt", "Service IDs active or removed by date", "Used to isolate trips running on the representative weekday."],
        ["feed_info.txt", "Feed validity dates", "Used as a data-validity check."],
    ]
    add_table(doc, gtfs_rows, ["GTFS table", "Main fields", "Pipeline role"], [1.2, 2.1, 3.2], font_size=8)

    doc.add_heading("3. Stage 1 - GTFS Exploration and Weekday Selection", level=1)
    add_paragraph(
        doc,
        "Objective: create a clean working subset of trips that actually operate on the selected representative weekday.",
    )
    add_numbered(
        doc,
        [
            "Load the GTFS ZIP and inspect required tables.",
            f"Confirm that the feed covers {FEED_PERIOD}.",
            f"Select {ANALYSIS_DATE} as the representative weekday to avoid mixing weekday and weekend service patterns.",
            "Use calendar_dates.txt to identify active service_id values where exception_type = 1 for the analysis date.",
            "Filter trips.txt to those active service IDs and join trips to routes.txt for readable route names.",
            "Export a route-level starting summary with scheduled trips, directions, and distinct shapes.",
        ],
    )
    add_callout(
        doc,
        "Why this matters",
        "Transit schedules vary by day type. Isolating one normal weekday keeps the score internally consistent and prevents school, weekend, or special-event patterns from distorting the route comparison.",
    )

    doc.add_heading("4. Stage 2 - Route Geometry and Stop Metrics", level=1)
    add_paragraph(
        doc,
        "Objective: describe the physical structure of each route using shape geometry and stops. This stage measures how direct each route is and how stop-heavy it is.",
    )
    add_bullets(
        doc,
        [
            "Shape path length: ordered shape points are converted into segment distances with a Haversine distance function and summed by shape_id.",
            "Route-level length: shape lengths are averaged using scheduled-trip weighting, so common branches count more than rare branches.",
            "Directness ratio: straight-line endpoint distance divided by route path distance. Higher values are more direct.",
            "Loop handling: routes with endpoint straight-line distance below 1 km are flagged as loop/circulator routes because directness can be mathematically misleading.",
            "Stop density: unique stops divided by average shape path kilometres. This is treated as an operating efficiency signal, not an access-quality judgement.",
        ],
    )
    geometry_summary = [
        ["Routes", len(geometry)],
        ["Average path length range", f"{geometry['avg_shape_path_km'].min():.2f} to {geometry['avg_shape_path_km'].max():.2f} km"],
        ["Directness range", f"{geometry['avg_directness_ratio'].min():.3f} to {geometry['avg_directness_ratio'].max():.3f}"],
        ["Stops per km range", f"{geometry['stops_per_km'].min():.2f} to {geometry['stops_per_km'].max():.2f}"],
    ]
    add_table(doc, geometry_summary, ["Metric", "Observed range/value"], [2.2, 4.3], font_size=9)

    doc.add_heading("5. Stage 3 - Scheduled Speed and Travel Time", level=1)
    add_paragraph(
        doc,
        "Objective: estimate how fast each route is scheduled to operate using trip durations and shape-based trip distances.",
    )
    add_bullets(
        doc,
        [
            "GTFS times are parsed manually into seconds after service-day start so values after 24:00:00 are handled correctly.",
            "Trip duration is calculated from the first scheduled departure to the last scheduled arrival on each trip.",
            "Because shape_dist_traveled is blank in this feed, trip distance is calculated from shape geometry.",
            "Trip-level scheduled speed is distance divided by duration, then summarized by route.",
            "Implausible records are kept in a diagnostic output and excluded from route-level speed averages to avoid distorted scores.",
        ],
    )
    speed_summary = [
        ["Routes with speed metrics", len(speed)],
        ["Average scheduled speed range", f"{speed['avg_scheduled_speed_kmh'].min():.2f} to {speed['avg_scheduled_speed_kmh'].max():.2f} km/h"],
        ["System average scheduled speed", f"{speed['avg_scheduled_speed_kmh'].mean():.2f} km/h"],
        ["Median route duration", f"{speed['median_duration_min'].median():.2f} minutes"],
    ]
    add_table(doc, speed_summary, ["Metric", "Value"], [2.4, 4.1], font_size=9)

    doc.add_heading("6. Stage 4 - Frequency, Availability, and Directional Balance", level=1)
    add_paragraph(
        doc,
        "Objective: measure how available each route is across the day, not only how fast it is. For riders, a route that is fast but infrequent may still be less useful.",
    )
    add_bullets(
        doc,
        [
            "Trips are grouped into practical service periods: early morning, AM peak, midday, PM peak, evening, and late night.",
            "Headways are calculated from consecutive departures within each route and direction.",
            "Route-level frequency uses median headway because it is less distorted by a few long late-night gaps.",
            "Service span measures the time from first departure to last arrival.",
            "Directional balance is the smaller directional trip count divided by the larger directional trip count; values close to 1 indicate balanced bidirectional service.",
            "The all-day flag identifies routes with meaningful service across the day, but it is retained as a diagnostic rather than directly included in the weighted final score.",
        ],
    )
    frequency_summary = [
        ["Total scheduled weekday trips", f"{int(frequency['scheduled_trips'].sum()):,}"],
        ["All-day routes", int(frequency["all_day_service_flag"].sum())],
        ["Median route headway", f"{frequency['median_headway_min'].median():.2f} minutes"],
        ["Service span range", f"{frequency['full_service_span_hours'].min():.2f} to {frequency['full_service_span_hours'].max():.2f} hours"],
    ]
    add_table(doc, frequency_summary, ["Metric", "Value"], [2.4, 4.1], font_size=9)

    doc.add_heading("7. Stage 5 - Final Efficiency Scoring", level=1)
    add_paragraph(
        doc,
        "Objective: convert different route metrics into comparable 0-100 component scores and combine them into a transparent final efficiency score.",
    )
    add_paragraph(
        doc,
        "Percentile scoring: each raw metric is ranked against other MiWay routes. For metrics where higher is better, a higher percentile gets a higher score. For metrics where lower is better, the percentile direction is reversed. Missing values receive a neutral score of 50.",
    )
    formula_rows = [
        ["Speed score", "Percentile(avg_scheduled_speed_kmh)", "Higher is better"],
        ["Headway score", "Reverse percentile(median_headway_min)", "Lower headway is better"],
        ["Trip volume score", "Percentile(scheduled_trips)", "Higher volume is better"],
        ["Frequency score", "0.70 x headway_score + 0.30 x trip_volume_score", "Balances wait time and service volume"],
        ["Service span score", "Percentile(full_service_span_hours)", "Longer span is better"],
        ["Directness score", "Percentile(directness_ratio_for_scoring)", "Higher directness is better; loops receive neutral 50"],
        ["Stop-spacing efficiency", "Reverse percentile(stops_per_km)", "Lower stop density is treated as more efficient"],
        ["Directional balance", "Percentile(directional_balance_ratio)", "More balanced service is better"],
    ]
    add_table(doc, formula_rows, ["Component", "Formula", "Interpretation"], [1.55, 2.75, 2.2], font_size=8)

    weights_rows = [
        ["Scheduled speed", "25%"],
        ["Frequency", "25%"],
        ["Service span", "15%"],
        ["Route directness", "15%"],
        ["Stop-spacing efficiency", "10%"],
        ["Directional balance", "10%"],
    ]
    add_table(doc, weights_rows, ["Final efficiency component", "Weight"], [3.0, 1.2], font_size=9)
    add_paragraph(
        doc,
        "Final efficiency score = 0.25(speed) + 0.25(frequency) + 0.15(service span) + 0.15(directness) + 0.10(stop spacing) + 0.10(directional balance).",
    )
    add_paragraph(
        doc,
        "Score bands: very low is below 40, low is 40 to 55, moderate is 55 to 70, high is 70 to 85, and very high is above 85.",
    )

    doc.add_heading("8. Advocacy Priority Lens", level=1)
    add_paragraph(
        doc,
        "Objective: identify routes where advocacy or planning review may be especially useful. This is separate from efficiency ranking because an inefficient route with many scheduled trips may represent a bigger opportunity than a low-volume special-purpose route.",
    )
    add_paragraph(
        doc,
        "Improvement need score = 100 - final_efficiency_score. Service importance score = percentile(scheduled_trips). Advocacy priority score = 0.60 x improvement_need_score + 0.40 x service_importance_score.",
    )
    add_df_table(
        doc,
        priority.head(10),
        [
            "advocacy_priority_rank",
            "route_short_name",
            "route_long_name",
            "advocacy_priority_score",
            "final_efficiency_score",
            "scheduled_trips",
        ],
        ["Rank", "Route", "Name", "Priority", "Efficiency", "Trips"],
        [0.55, 0.6, 2.35, 0.8, 0.8, 0.65],
        font_size=8,
    )

    doc.add_heading("9. ArcGIS Handoff and Map Production", level=1)
    add_paragraph(
        doc,
        "Objective: connect the scoring table back to route geometry so the results can be mapped and interpreted spatially in ArcGIS Pro.",
    )
    add_numbered(
        doc,
        [
            "Read GTFS shapes.txt and build ordered LineString geometries from shape point sequences.",
            "Create route lookup fields from routes.txt and shape-to-route assignments from trips.txt.",
            "Join final route efficiency fields from final_route_efficiency_ranking_2026-05-05.csv.",
            "Export one detailed GeoJSON with every route/direction/shape combination.",
            "Export one representative GeoJSON using the longest/highest-trip shape per route direction for a cleaner overview map.",
            "Load the representative route-line layer into ArcGIS Pro and symbolize by final_efficiency_score or advocacy_priority_score.",
        ],
    )
    arcgis_rows = [
        ["Outputs/ArcGIS/miway_route_shape_lines.geojson", "Detailed route-shape line features; useful for troubleshooting branches and variants."],
        ["Outputs/ArcGIS/miway_representative_route_lines.geojson", "Cleaner route-direction overview layer used for the main maps."],
        ["final_efficiency_score", "Graduated line color for the efficiency map."],
        ["advocacy_priority_score", "Graduated line color for the advocacy priority map."],
        ["route_label", "Suggested label field for map annotation."],
    ]
    add_table(doc, arcgis_rows, ["Layer/field", "Use"], [2.8, 3.7], font_size=8.5)

    doc.add_heading("10. Main Outputs", level=1)
    output_rows = [
        ["route_geometry_stop_metrics_2026-05-05.csv", "Route geometry, directness, stop counts, and stop density."],
        ["route_scheduled_speed_metrics_2026-05-05.csv", "Scheduled speed, trip duration, and speed spread metrics."],
        ["route_frequency_service_metrics_2026-05-05.csv", "Service span, headways, period service, directional balance, and all-day flag."],
        ["route_efficiency_score_components_2026-05-05.csv", "Full joined metric table with raw metrics and component scores."],
        ["final_route_efficiency_ranking_2026-05-05.csv", "Final ranked route table used for reporting and ArcGIS joins."],
        ["advocacy_priority_routes_2026-05-05.csv", "Top advocacy-review candidates and supporting metrics."],
        ["miway_representative_route_lines.geojson", "ArcGIS-ready line layer for the main maps."],
    ]
    add_table(doc, output_rows, ["Output file", "Purpose"], [2.75, 3.75], font_size=8)

    doc.add_heading("11. Key Results Snapshot", level=1)
    snapshot_rows = [
        ["Routes analyzed", len(ranking)],
        ["Final efficiency score range", f"{ranking['final_efficiency_score'].min():.2f} to {ranking['final_efficiency_score'].max():.2f}"],
        ["Median final efficiency score", f"{ranking['final_efficiency_score'].median():.2f}"],
        ["Highest scoring route", f"{ranking.sort_values('efficiency_rank').iloc[0]['route_short_name']} - {ranking.sort_values('efficiency_rank').iloc[0]['route_long_name']}"],
        ["Top advocacy priority route", f"{priority.iloc[0]['route_short_name']} - {priority.iloc[0]['route_long_name']}"],
    ]
    add_table(doc, snapshot_rows, ["Result", "Value"], [2.4, 4.1], font_size=9)

    doc.add_heading("12. Quality Checks", level=1)
    add_bullets(
        doc,
        [
            "Confirm all required GTFS files exist before running the notebooks.",
            "Confirm the selected analysis date is inside the feed validity period.",
            "Check that every active trip has a route_id and shape_id before geometry calculations.",
            "Keep diagnostics for implausible trip speeds rather than silently dropping records.",
            "Confirm final ranking has one row per active route and no missing final efficiency score.",
            "Review loop/circulator route flags before interpreting directness results.",
            "Inspect ArcGIS joins by checking that route IDs, scores, and score bands appear on route-line features.",
        ],
    )

    doc.add_heading("13. Interpretation Rules and Limitations", level=1)
    add_bullets(
        doc,
        [
            "The results use scheduled GTFS data, not real-time travel time, on-time performance, crowding, or rider counts.",
            "A low score is a starting point for investigation, not evidence that a route should be cut.",
            "Stop spacing is an efficiency signal but can conflict with accessibility, walk distance, and equity.",
            "School, loop, and limited-purpose routes should be interpreted separately from all-day grid routes.",
            "The scoring weights are transparent and defensible for a first pass, but they should be sensitivity-tested before formal policy use.",
            "Advocacy priority is meant to direct further research toward high-need routes, not replace community consultation.",
        ],
    )

    doc.add_heading("14. Reproduction Checklist", level=1)
    add_numbered(
        doc,
        [
            "Place the latest MiWay GTFS feed at Data/google_transit.zip and extract it to Data/google_transit.",
            "Run notebooks 01 through 05 in order using the same representative service day.",
            "Review exported CSVs in Outputs for row counts, missing values, and obvious outliers.",
            "Run Scripts/create_arcgis_route_lines.py from the project root to rebuild GeoJSON route-line layers.",
            "Open ArcGIS Pro and add Outputs/ArcGIS/miway_representative_route_lines.geojson.",
            "Symbolize final_efficiency_score for the efficiency map and advocacy_priority_score for the advocacy map.",
            "Export final layouts to Reports with legends, scale bar, north arrow, and data-source attribution.",
        ],
    )

    doc.add_heading("15. Final Map Outputs", level=1)
    add_image_if_exists(
        doc,
        REPORT_DIR / "MiWay Route Efficiency Score2.png",
        "Final map output: route efficiency score across the MiWay network.",
    )
    add_image_if_exists(
        doc,
        REPORT_DIR / "MiWay Top 10 Advocacy Priority Routes2.png",
        "Final map output: top advocacy priority routes derived from the scoring framework.",
    )

    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_heading("Appendix A - Top Efficiency Routes", level=1)
    add_df_table(
        doc,
        ranking.sort_values("efficiency_rank").head(10),
        [
            "efficiency_rank",
            "route_short_name",
            "route_long_name",
            "final_efficiency_score",
            "scheduled_trips",
            "avg_scheduled_speed_kmh",
            "median_headway_min",
        ],
        ["Rank", "Route", "Name", "Score", "Trips", "Speed", "Headway"],
        [0.55, 0.6, 2.1, 0.75, 0.65, 0.8, 0.8],
        font_size=8,
    )

    doc.add_heading("Appendix B - Notebook and Script Inventory", level=1)
    inventory_rows = [
        ["Notebooks/01_miway_gtfs_exploration.ipynb", "Data loading, feed validation, active weekday trips."],
        ["Notebooks/02_miway_route_geometry_metrics.ipynb", "Geometry, directness, and stop-density metrics."],
        ["Notebooks/03_miway_scheduled_speed_travel_time.ipynb", "Trip duration and scheduled speed metrics."],
        ["Notebooks/04_miway_frequency_service_availability.ipynb", "Headways, service span, period service, and directional balance."],
        ["Notebooks/05_miway_final_route_efficiency_ranking.ipynb", "Component scores, final ranking, and advocacy priority score."],
        ["Scripts/create_arcgis_route_lines.py", "ArcGIS-ready route-line GeoJSON generation."],
    ]
    add_table(doc, inventory_rows, ["File", "Role"], [3.15, 3.35], font_size=8)

    doc.core_properties.title = "MiWay Route Efficiency and Advocacy Priority Analysis Pipeline"
    doc.core_properties.subject = "GTFS analysis workflow and ArcGIS handoff documentation"
    doc.core_properties.author = "Raiyan Bin Rizwan"
    doc.save(DOC_PATH)
    return DOC_PATH


if __name__ == "__main__":
    print(build_doc())
