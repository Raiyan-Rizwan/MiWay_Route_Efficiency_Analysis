from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement


SRC = Path(r"C:\Users\Raiyan Rizwan\Desktop\UTM Transit Housing Advocacy\Raiyan_Bin_Rizwan_Resume_Updated.docx")
OUT = Path(r"C:\Users\Raiyan Rizwan\Desktop\UTM Transit Housing Advocacy\Miway Route Efficiency\Raiyan_Bin_Rizwan_PepsiCo_AI_ML_Scientist_Resume.docx")


def clear_runs(paragraph):
    for run in list(paragraph.runs):
        run._element.getparent().remove(run._element)


def set_text_like(paragraph, text):
    bold = paragraph.runs[0].bold if paragraph.runs else None
    italic = paragraph.runs[0].italic if paragraph.runs else None
    clear_runs(paragraph)
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic


def set_heading(paragraph, text):
    clear_runs(paragraph)
    run = paragraph.add_run(text)
    run.bold = True


def set_role(paragraph, title, date):
    clear_runs(paragraph)
    run = paragraph.add_run(title)
    run.bold = True
    paragraph.add_run("\t" + date)


def set_italic(paragraph, text):
    clear_runs(paragraph)
    run = paragraph.add_run(text)
    run.italic = True


def set_label(paragraph, label, text):
    clear_runs(paragraph)
    run = paragraph.add_run(label + ": ")
    run.bold = True
    paragraph.add_run(text)


def insert_after(paragraph, text="", style=None, clone_ppr_from=None):
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_para = paragraph._parent.add_paragraph()
    new_para._p.getparent().remove(new_para._p)
    new_p.addnext(new_para._p)
    new_para._p.getparent().remove(new_p)
    if clone_ppr_from is not None and clone_ppr_from._p.pPr is not None:
        new_para._p.insert(0, deepcopy(clone_ppr_from._p.pPr))
    if style is not None:
        new_para.style = style
    if text:
        new_para.add_run(text)
    return new_para


def delete_paragraph(paragraph):
    p = paragraph._element
    p.getparent().remove(p)
    paragraph._p = paragraph._element = None


doc = Document(SRC)
p = doc.paragraphs

# Add a compact role-focused summary after contact information.
summary_heading = insert_after(p[1], "", clone_ppr_from=p[2])
set_heading(summary_heading, "SUMMARY")
summary_body = insert_after(
    summary_heading,
    "AI/ML new graduate with applied experience building production-oriented speech AI pipelines, SQL-backed analytics tools, and statistical models on large datasets. Strong fit for business-embedded AI work: translating user needs into technical requirements, evaluating model performance, and communicating insights for adoption.",
    clone_ppr_from=p[4],
)

# Re-read after insertion because paragraph positions changed.
p = doc.paragraphs

replacements = {
    "Relevant coursework: Time Series Analysis, Multivariate Methods, Bayesian, Software Design, Linear Algebra, Real Analysis, Stochastic Processes, Supervised & Unsupervised ML Models, Web App Development":
        "Relevant coursework: Statistical Modeling, Supervised & Unsupervised ML, Time Series Analysis, Multivariate Methods, Bayesian Inference, Software Design, Linear Algebra, Stochastic Processes, Web App Development",

    "Architected an end-to-end Azure Custom Speech training pipeline in Node.js, enabling automatic acoustic model fine-tuning on clinical audio data — one of few student-built systems integrating real clinical speech AI at this depth":
        "Delivered an end-to-end Azure Custom Speech proof of concept in Node.js, translating clinical speech-recognition needs into a technical pipeline for automated acoustic model fine-tuning on real-world audio data",
    "Designed a fully automated retraining workflow using onboarding voice recordings and word-bank entries to continuously fine-tune a Whisper Large V2–based model via Azure Speech REST API v3.2, improving transcription accuracy without manual intervention":
        "Designed an automated retraining workflow using onboarding voice recordings and word-bank entries to fine-tune a Whisper Large V2-based model via Azure Speech REST API v3.2, reducing manual model maintenance",
    "Built robust audio preprocessing (PCM WAV conversion via ffmpeg, deduplication, 80/20 train/test splitting, zip packaging) and automated dataset registration to Azure Blob Storage with SAS URL generation":
        "Built repeatable data-preparation steps for large audio datasets, including PCM WAV conversion, deduplication, 80/20 train/test splitting, zip packaging, Azure Blob registration, and SAS URL generation",
    "Integrated the ML pipeline into an Express.js backend using a decoupled Node.js polling worker on Oracle Cloud (PM2), bypassing serverless timeout constraints with async job queuing via Supabase":
        "Integrated the ML workflow into an Express.js backend with a decoupled Node.js polling worker on Oracle Cloud (PM2), using Supabase-backed asynchronous job queues to support reliable business-process adoption",
    "Evaluated model performance using Word Error Rate (WER) metrics against a held-out test set; researched Azure Custom Speech domain-specific medical transcription boosting strategies":
        "Evaluated model performance with Word Error Rate (WER) on held-out test data and communicated domain-specific transcription improvement strategies to technical stakeholders",

    "Visualized protein structures and drug binding sites using Python and Amazon AWS cloud computing services":
        "Analyzed and visualized protein structures and drug binding sites using Python and AWS, helping convert complex scientific data into usable research insights",
    "Built HTML-based web interfaces integrated with SQL queries for protein structure data analysis and visualization":
        "Built SQL-driven HTML interfaces for protein-structure data exploration, enabling non-technical users to query, filter, and interpret relational research datasets",
    "Wrote Python (Pandas) and SQL scripts to generate Excel-based visualizations for biological data pipelines.":
        "Wrote Python, Pandas, and SQL scripts to clean biological datasets and generate Excel-based visualizations for recurring analysis workflows",
    "Developed frontend components to integrate visualizations into a user-friendly interface, improving accessibility of analysis tools":
        "Collaborated across technical and research needs to integrate visualizations into a user-friendly interface, improving accessibility and adoption of analysis tools",

    "Provide exceptional customer service by assisting guests, adults and children, with arcade game instructions and ensuring a positive gaming experience across the facility":
        "Communicate instructions and safety protocols clearly to diverse guests in a fast-paced environment, adapting explanations for children, adults, and groups",
    "Facilitate and lead group activities including archery and axe throwing, clearly communicating safety protocols, game rules, and proper techniques":
        "Identify operational issues, coordinate with technicians, and maintain service continuity during high-traffic periods",

    "Built an end-to-end automated trading pipeline on all 500+ S&P 500 constituents (2018–2026), engineering 6 technical features per stock (Garman-Klass Volatility, RSI, Bollinger Bands, ATR, MACD, Dollar Volume) from raw OHLCV data via yfinance.":
        "Built an end-to-end ML analytics pipeline on 500+ S&P 500 constituents (2018-2026), engineering technical features per stock from raw OHLCV data to support monthly decision-making",
    "Estimated each stock’s exposure to Fama-French 5 risk factors using rolling OLS regression (statsmodels), then applied K-Means clustering (k=4, RSI-anchored initialization) monthly to group stocks into behavioural regimes and select the high-momentum cluster.":
        "Estimated stock exposure to Fama-French 5 risk factors with rolling OLS regression, then applied K-Means clustering monthly to identify behavioural regimes and select high-momentum candidates",
    "Optimized monthly portfolio weights via Efficient Frontier (max Sharpe Ratio, 10% single-stock cap) using PyPortfolioOpt, with a 5-year rolling average dollar volume filter retaining only the top 150 most liquid stocks per period.":
        "Optimized monthly portfolio weights via Efficient Frontier using PyPortfolioOpt, applying liquidity filters and cap constraints to convert model output into an actionable allocation strategy",

    "Analysed 19,364 respondents across 10 countries from WVS Wave 7 using PCA, factor analysis (varimax), GMM clustering, and LDA to identify three latent value dimensions: traditional engagement, subjective well-being, and social openness.":
        "Analyzed 19,364 respondents across 10 countries from WVS Wave 7 using PCA, factor analysis, GMM clustering, and LDA to identify interpretable latent value dimensions",
    "Handled mixed-scale ordinal survey data (binary, 4-pt, 5-pt, 10-pt), applied listwise deletion retaining 85.4% of sample, and standardized all variables before analysis.":
        "Prepared mixed-scale ordinal survey data, retained 85.4% of the sample after quality filtering, and standardized variables before statistical modeling",
    "Validated results via Minimum Covariance Determinant (MCD) robust PCA (congruence ≥ 0.91), bootstrap eigenvalue CIs, and GMM stability checks (ARI = 0.90 across 20 splits).":
        "Validated findings using MCD robust PCA, bootstrap eigenvalue confidence intervals, and GMM stability checks, strengthening confidence in model-driven insights",
    "LDA achieved 36.2% cross-validated accuracy (vs 10% chance), confirming systematic but partial between-country value differences; within-country variation substantially exceeded between-country differences.":
        "Translated LDA validation results into a clear interpretation: country-level differences were systematic but partial, while within-country variation remained the dominant pattern",
}

for para in doc.paragraphs:
    text = para.text.strip()
    if text in replacements:
        set_text_like(para, replacements[text])

# Compress less-relevant customer-service experience by removing redundant bullets.
for text in [
    "While leading activities, ensure compliance with safety guidelines and maintain order among guests.",
    "Promptly identify and report any arcade machine issues to technicians, minimizing downtime and maintaining operational efficiency",
    "Maintain cleanliness and organization of the games floor, proactively addressing customers' immediate needs",
    "Accurately handle cash transactions and manage inventory, including restocking prizes and goods to ensure optimal availability and presentation",
]:
    for para in list(doc.paragraphs):
        if para.text.strip() == text:
            delete_paragraph(para)

# Remove the inconsistent duplicate project block to keep the resume focused.
remove_block = False
for para in list(doc.paragraphs):
    text = para.text.strip()
    if text.startswith("Affect of Covid-19 on Canadian Businesses"):
        remove_block = True
    if remove_block:
        delete_paragraph(para)
        if text == "Designed SQL queries and Python scripts to power interactive HTML dashboards, enabling non-technical researchers to explore structural data.":
            remove_block = False

# Refresh technical skills for the PepsiCo AI/ML Scientist posting.
for para in doc.paragraphs:
    text = para.text.strip()
    if text.startswith("Languages: "):
        set_label(para, "Languages", "Python, R, SQL, JavaScript (Node.js), Bash")
    elif text.startswith("ML / Stats: "):
        set_label(para, "ML / Stats", "Statistical modeling, supervised and unsupervised learning, clustering, regression, time series, PCA, factor analysis, GMM, LDA, robust estimation, Bayesian inference")
    elif text.startswith("Libraries: "):
        set_label(para, "Libraries", "Pandas, NumPy, scikit-learn, statsmodels, Matplotlib, Seaborn, PyPortfolioOpt")
    elif text.startswith("Cloud & Infra: "):
        set_label(para, "AI / Cloud", "Azure Custom Speech, Azure Blob Storage, Whisper Large V2, AWS, Oracle Cloud, Supabase, PM2")
    elif text.startswith("Tools: "):
        set_label(para, "Data / Visualization", "SQL, relational data workflows, Excel dashboards, HTML dashboards, Git, ffmpeg, Express.js")
    elif text.startswith("Mathematics: "):
        set_label(para, "Business Translation", "Technical requirement translation, stakeholder communication, model evaluation, insight storytelling")

# Normalize headings/capitalization.
for para in doc.paragraphs:
    if para.text.strip() == "Awards":
        set_heading(para, "AWARDS")

doc.save(OUT)
print(OUT)
