from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Inches, Pt


PROJECT_ROOT = Path(r"C:\Users\Raiyan Rizwan\Desktop\UTM Transit Housing Advocacy\Miway Route Efficiency")
OUT_PATH = PROJECT_ROOT / "Reports" / "ArcGIS_Pro_GTFS_Route_Efficiency_Pipeline.docx"


CODE_BLOCKS = [
    (
        "1. Set up paths and geodatabase",
        r'''import arcpy
import pandas as pd
from pathlib import Path

project_root = Path(r"C:\Users\Raiyan Rizwan\Desktop\UTM Transit Housing Advocacy\Miway Route Efficiency")

data_dir = project_root / "Data" / "google_transit"
outputs_dir = project_root / "Outputs"

stops_csv = outputs_dir / "arcgis_stops_clean.csv"
shapes_csv = outputs_dir / "arcgis_shapes_clean.csv"
ranking_csv = outputs_dir / "final_route_efficiency_ranking_2026-05-05.csv"

gdb = project_root / "Outputs" / "MiWay_Route_Efficiency.gdb"

if not arcpy.Exists(str(gdb)):
    arcpy.management.CreateFileGDB(str(gdb.parent), gdb.name)

arcpy.env.workspace = str(gdb)
arcpy.env.overwriteOutput = True

wgs84 = arcpy.SpatialReference(4326)
utm17 = arcpy.SpatialReference(26917)  # NAD83 / UTM zone 17N, good for Mississauga''',
    ),
    (
        "2. Create stops point layer",
        r'''stops_points_wgs84 = str(gdb / "miway_stops_wgs84")
stops_points_utm = str(gdb / "miway_stops_utm17")

arcpy.management.XYTableToPoint(
    in_table=str(stops_csv),
    out_feature_class=stops_points_wgs84,
    x_field="stop_lon",
    y_field="stop_lat",
    coordinate_system=wgs84
)

arcpy.management.Project(
    in_dataset=stops_points_wgs84,
    out_dataset=stops_points_utm,
    out_coor_system=utm17
)

print("Created stops:", stops_points_utm)''',
    ),
    (
        "3. Create route shape point layer",
        r'''shape_points_wgs84 = str(gdb / "miway_shape_points_wgs84")
shape_points_utm = str(gdb / "miway_shape_points_utm17")

arcpy.management.XYTableToPoint(
    in_table=str(shapes_csv),
    out_feature_class=shape_points_wgs84,
    x_field="shape_pt_lon",
    y_field="shape_pt_lat",
    coordinate_system=wgs84
)

arcpy.management.Project(
    in_dataset=shape_points_wgs84,
    out_dataset=shape_points_utm,
    out_coor_system=utm17
)

print("Created shape points:", shape_points_utm)''',
    ),
    (
        "4. Convert shape points into route lines",
        r'''shape_lines = str(gdb / "miway_route_shapes_utm17")

arcpy.management.PointsToLine(
    Input_Features=shape_points_utm,
    Output_Feature_Class=shape_lines,
    Line_Field="shape_id_text",
    Sort_Field="shape_pt_sequence"
)

print("Created route shape lines:", shape_lines)''',
    ),
    (
        "5. Create shape-to-route lookup table",
        r'''trips = pd.read_csv(data_dir / "trips.txt", dtype=str)
routes = pd.read_csv(data_dir / "routes.txt", dtype=str)

shape_route_lookup = (
    trips[["shape_id", "route_id"]]
    .dropna()
    .drop_duplicates()
    .merge(
        routes[["route_id", "route_short_name", "route_long_name"]],
        on="route_id",
        how="left"
    )
)

shape_route_lookup["shape_id_text"] = shape_route_lookup["shape_id"].astype(str)

lookup_csv = outputs_dir / "arcgis_shape_route_lookup.csv"
shape_route_lookup.to_csv(lookup_csv, index=False)

print("Saved:", lookup_csv)''',
    ),
    (
        "6. Import lookup and ranking tables into the geodatabase",
        r'''lookup_table = str(gdb / "shape_route_lookup")
ranking_table = str(gdb / "route_efficiency_ranking")

arcpy.conversion.TableToTable(str(lookup_csv), str(gdb), "shape_route_lookup")
arcpy.conversion.TableToTable(str(ranking_csv), str(gdb), "route_efficiency_ranking")

print("Imported lookup and ranking tables")''',
    ),
    (
        "7. Join lookup to route shape lines",
        r'''shape_lines_with_route = str(gdb / "miway_route_shapes_with_route")

arcpy.management.MakeFeatureLayer(shape_lines, "shape_lines_lyr")

arcpy.management.AddJoin(
    in_layer_or_view="shape_lines_lyr",
    in_field="shape_id_text",
    join_table=lookup_table,
    join_field="shape_id_text",
    join_type="KEEP_COMMON"
)

arcpy.management.CopyFeatures("shape_lines_lyr", shape_lines_with_route)

print("Created:", shape_lines_with_route)''',
    ),
    (
        "8. Check route ID field name before final join",
        r'''arcpy.management.MakeFeatureLayer(shape_lines_with_route, "routes_with_ids_lyr")

for f in arcpy.ListFields(shape_lines_with_route):
    print(f.name)''',
    ),
    (
        "9. Join route efficiency metrics onto route lines",
        r'''final_routes = str(gdb / "miway_route_efficiency_lines")

# Adjust this if the field printed in the previous step has a slightly different name.
route_id_field = "shape_route_lookup_route_id"

arcpy.management.AddJoin(
    in_layer_or_view="routes_with_ids_lyr",
    in_field=route_id_field,
    join_table=ranking_table,
    join_field="route_id",
    join_type="KEEP_COMMON"
)

arcpy.management.CopyFeatures("routes_with_ids_lyr", final_routes)

print("Final route efficiency layer:", final_routes)''',
    ),
    (
        "10. Add final layers to the current ArcGIS Pro map",
        r'''aprx = arcpy.mp.ArcGISProject("CURRENT")
m = aprx.activeMap

m.addDataFromPath(final_routes)
m.addDataFromPath(stops_points_utm)

aprx.save()

print("Added final routes and stops to current ArcGIS Pro map.")''',
    ),
]


def add_code_block(document: Document, code: str) -> None:
    style = document.styles["Normal"]
    for line in code.splitlines():
        p = document.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.2)
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)


def build_doc() -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10.5)
    styles["Title"].font.name = "Arial"
    styles["Title"].font.size = Pt(18)
    styles["Heading 1"].font.name = "Arial"
    styles["Heading 1"].font.size = Pt(13)

    doc.add_heading("ArcGIS Pro GTFS Route Efficiency Pipeline", 0)
    doc.add_paragraph(
        "Run these blocks in order after the Create GTFS notebook has produced "
        "arcgis_stops_clean.csv, arcgis_shapes_clean.csv, and the final route ranking CSV."
    )
    doc.add_paragraph(
        "Note: Step 8 prints the joined field names. If the route ID field name differs from "
        "shape_route_lookup_route_id, update route_id_field in Step 9 before running it."
    )

    for title, code in CODE_BLOCKS:
        doc.add_heading(title, level=1)
        add_code_block(doc, code)

    doc.add_heading("Recommended map outputs", level=1)
    for item in [
        "Final Efficiency Ranking Map: symbolize by final_efficiency_score.",
        "Advocacy Priority Map: symbolize by advocacy_priority_score.",
        "Score Band Map: symbolize by score_band.",
        "Scheduled Speed Map: symbolize by avg_scheduled_speed_kmh.",
        "Frequency Map: symbolize by median_headway_min.",
        "Directness Map: symbolize by avg_directness_ratio.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_paragraph("Main final feature class:")
    add_code_block(doc, r"Outputs\MiWay_Route_Efficiency.gdb\miway_route_efficiency_lines")

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build_doc()
