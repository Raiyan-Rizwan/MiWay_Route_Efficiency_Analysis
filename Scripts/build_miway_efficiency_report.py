from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = PROJECT_ROOT / "Outputs"
REPORT_DIR = PROJECT_ROOT / "Reports"
REPORT_DIR.mkdir(exist_ok=True)

REPORT_PATH = REPORT_DIR / "MiWay_Route_Efficiency_Advocacy_Report.docx"
SAFE_REPORT_PATH = REPORT_DIR / "MiWay_Route_Efficiency_Advocacy_Report_FIXED.docx"

ANALYSIS_DATE = "Tuesday, May 5, 2026"
FEED_PERIOD = "May 1, 2026 to June 28, 2026"


def load_data():
    ranking = pd.read_csv(OUTPUT_DIR / "final_route_efficiency_ranking_2026-05-05.csv")
    priority = pd.read_csv(OUTPUT_DIR / "advocacy_priority_routes_2026-05-05.csv")
    geometry = pd.read_csv(OUTPUT_DIR / "route_geometry_stop_metrics_2026-05-05.csv")
    speed = pd.read_csv(OUTPUT_DIR / "route_scheduled_speed_metrics_2026-05-05.csv")
    frequency = pd.read_csv(OUTPUT_DIR / "route_frequency_service_metrics_2026-05-05.csv")
    return ranking, priority, geometry, speed, frequency


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9E2EC", size="4"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
            tc_pr = row.cells[idx]._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 1440)))
            tc_w.set(qn("w:type"), "dxa")


def set_run_font(run, size=None, color=None, bold=None, italic=None, name="Calibri"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(22)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("0B2545")
    title.paragraph_format.space_after = Pt(6)

    for style_name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("MiWay Route Efficiency Analysis")
    set_run_font(run, size=9, color="666666")


def add_title_block(doc):
    p = doc.add_paragraph()
    p.style = "Title"
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.add_run("MiWay Route Efficiency Analysis")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run("Professional summary of GTFS-based route metrics and advocacy priorities")
    set_run_font(run, size=13, color="555555")

    meta_rows = [
        ("Prepared for", "UTM Transit Housing Advocacy"),
        ("Dataset", f"MiWay GTFS scheduled service feed, valid {FEED_PERIOD}"),
        ("Representative day", ANALYSIS_DATE),
        ("Analysis scope", "Scheduled service, route geometry, speed, frequency, and route-level scoring"),
    ]
    table = doc.add_table(rows=len(meta_rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_table_width(table, [1.45, 5.05])
    for row_idx, (label, value) in enumerate(meta_rows):
        for cell in table.rows[row_idx].cells:
            set_cell_margins(cell)
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        label_cell = table.rows[row_idx].cells[0]
        value_cell = table.rows[row_idx].cells[1]
        set_cell_shading(label_cell, "F2F4F7")
        label_run = label_cell.paragraphs[0].add_run(label)
        set_run_font(label_run, size=9.5, color="0B2545", bold=True)
        value_run = value_cell.paragraphs[0].add_run(value)
        set_run_font(value_run, size=9.5, color="111111")

    doc.add_paragraph()


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_table_width(table, [6.5])
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F6F9")
    set_cell_border(cell, color="B7C9DB", size="6")
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(title)
    set_run_font(run, size=11, color="0B2545", bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(0)
    run = p2.add_run(body)
    set_run_font(run, size=10.5, color="111111")
    doc.add_paragraph()


def add_metric_strip(doc, metrics):
    table = doc.add_table(rows=2, cols=len(metrics))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    widths = [6.5 / len(metrics)] * len(metrics)
    set_table_width(table, widths)
    for col, (label, value) in enumerate(metrics):
        label_cell = table.cell(0, col)
        value_cell = table.cell(1, col)
        for cell in [label_cell, value_cell]:
            set_cell_margins(cell, top=100, bottom=100, start=90, end=90)
            set_cell_border(cell, color="D9E2EC")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_shading(label_cell, "E8EEF5")
        set_cell_shading(value_cell, "FFFFFF")
        label_run = label_cell.paragraphs[0].add_run(label)
        set_run_font(label_run, size=8.5, color="1F4D78", bold=True)
        value_run = value_cell.paragraphs[0].add_run(str(value))
        set_run_font(value_run, size=15, color="0B2545", bold=True)
    doc.add_paragraph()


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_run_font(run, size=10.5, color="111111")


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_run_font(run, size=10.5, color="111111")


def add_table_from_df(doc, df, columns, headers, widths, font_size=8.5):
    table = doc.add_table(rows=1, cols=len(columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    set_table_width(table, widths)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, "F2F4F7")
        set_cell_margins(cell, top=90, bottom=90, start=90, end=90)
        set_cell_border(cell)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        set_run_font(run, size=font_size, color="0B2545", bold=True)

    for _, row in df.iterrows():
        cells = table.add_row().cells
        for idx, col in enumerate(columns):
            value = row[col]
            if pd.isna(value):
                text = "n/a"
            elif isinstance(value, float):
                text = f"{value:.2f}"
            else:
                text = str(value)
            cell = cells[idx]
            set_cell_margins(cell, top=80, bottom=80, start=90, end=90)
            set_cell_border(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            if col in {"efficiency_rank", "route_short_name", "scheduled_trips", "advocacy_priority_rank"}:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run(text)
            set_run_font(run, size=font_size, color="111111")
    doc.add_paragraph()
    return table


def add_section_heading(doc, text):
    doc.add_heading(text, level=1)


def add_paragraph(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=11, color="111111")


def build_report(report_path=REPORT_PATH):
    ranking, priority, geometry, speed, frequency = load_data()

    summary = {
        "routes": len(ranking),
        "total_trips": int(frequency["scheduled_trips"].sum()),
        "all_day": int(frequency["all_day_service_flag"].sum()),
        "score_range": f"{ranking['final_efficiency_score'].min():.2f}-{ranking['final_efficiency_score'].max():.2f}",
        "median_headway": f"{frequency['median_headway_min'].median():.1f} min",
        "avg_speed": f"{speed['avg_scheduled_speed_kmh'].mean():.1f} km/h",
    }

    doc = Document()
    style_document(doc)
    add_title_block(doc)

    add_callout(
        doc,
        "Bottom line",
        "This analysis converts MiWay's scheduled GTFS feed into a transparent route efficiency ranking and an advocacy priority list. "
        "The strongest routes combine high scheduled speed, frequent service, long service span, and balanced bidirectional service. "
        "The priority list highlights routes where improvement could matter because the route has meaningful service volume but weaker efficiency indicators.",
    )

    add_metric_strip(
        doc,
        [
            ("Routes analyzed", summary["routes"]),
            ("Weekday trips", f"{summary['total_trips']:,}"),
            ("All-day routes", summary["all_day"]),
            ("Score range", summary["score_range"]),
        ],
    )

    add_section_heading(doc, "Executive Summary")
    add_paragraph(
        doc,
        f"The project analyzed MiWay scheduled service using a GTFS feed valid from {FEED_PERIOD}. "
        f"To avoid mixing weekday and weekend patterns, the analysis uses {ANALYSIS_DATE} as the representative service day."
    )
    add_paragraph(
        doc,
        "Across 66 active routes, the final route efficiency score ranges from "
        f"{ranking['final_efficiency_score'].min():.2f} to {ranking['final_efficiency_score'].max():.2f}, "
        f"with a median score of {ranking['final_efficiency_score'].median():.2f}. "
        f"The system-level average scheduled route speed is {speed['avg_scheduled_speed_kmh'].mean():.2f} km/h, "
        f"and the median route-level scheduled headway is {frequency['median_headway_min'].median():.2f} minutes."
    )
    add_bullets(
        doc,
        [
            "The highest-scoring routes are not simply the routes with the most trips. They tend to combine speed, all-day coverage, relatively frequent service, and balanced directions.",
            "The advocacy priority lens surfaces routes that may deserve operational attention because they have substantial service or rider relevance but weaker efficiency indicators.",
            "This ranking should be treated as a screening tool. It does not replace ridership, on-time performance, crowding, equity, or rider experience data.",
        ],
    )

    add_section_heading(doc, "What the Five Notebooks Produced")
    notebook_rows = pd.DataFrame(
        [
            ["1", "GTFS exploration", "Loaded source tables, confirmed feed dates, selected Tuesday, May 5, 2026, and isolated active weekday trips."],
            ["2", "Route geometry metrics", "Calculated route length, directness, unique stops, average stops per trip, and stops per kilometre."],
            ["3", "Scheduled speed and travel time", "Parsed GTFS times beyond midnight, calculated trip duration, distance, and scheduled speed by route."],
            ["4", "Frequency and availability", "Calculated first and last departures, service span, headways, period-level trips, all-day service, and directional balance."],
            ["5", "Final efficiency ranking", "Combined geometry, speed, and service metrics into final scores, ranks, score bands, and advocacy priority scores."],
        ],
        columns=["Notebook", "Focus", "Output"],
    )
    add_table_from_df(
        doc,
        notebook_rows,
        ["Notebook", "Focus", "Output"],
        ["Notebook", "Focus", "What it added"],
        [0.7, 1.65, 4.15],
        font_size=8.5,
    )

    add_section_heading(doc, "Scoring Method")
    add_paragraph(
        doc,
        "The final score uses percentile-based component scores, which means each route is compared against other MiWay routes in the same feed. "
        "Higher scores represent stronger scheduled performance relative to the rest of the system on the selected weekday."
    )
    scoring_rows = pd.DataFrame(
        [
            ["Scheduled speed", "25%", "Higher average scheduled speed scores better."],
            ["Frequency", "25%", "Lower median headway and higher trip volume score better."],
            ["Service span", "15%", "Longer daily span from first departure to last arrival scores better."],
            ["Route directness", "15%", "More direct non-loop route geometry scores better."],
            ["Stop-spacing efficiency", "10%", "Fewer stops per kilometre scores better as an operating efficiency signal."],
            ["Directional balance", "10%", "More even service between directions scores better."],
        ],
        columns=["Component", "Weight", "Interpretation"],
    )
    add_table_from_df(
        doc,
        scoring_rows,
        ["Component", "Weight", "Interpretation"],
        ["Component", "Weight", "Interpretation"],
        [1.55, 0.8, 4.15],
        font_size=8.5,
    )
    add_callout(
        doc,
        "Important caveat",
        "Stop density is an operating efficiency signal, not a service quality verdict. More stops can slow a route, but they can also improve access for riders. "
        "Similarly, loop routes can appear geometrically inefficient because their endpoints are close together; the notebook neutralizes directness for near-loop routes to avoid penalizing that design artifact.",
    )

    doc.add_page_break()
    add_section_heading(doc, "Final Route Efficiency Ranking")
    add_paragraph(
        doc,
        "The table below shows the ten highest-scoring routes in the final ranking. These routes generally combine strong speed, frequent service, long span, and balanced directions."
    )
    top = ranking.head(10).copy()
    add_table_from_df(
        doc,
        top,
        [
            "efficiency_rank",
            "route_short_name",
            "route_long_name",
            "final_efficiency_score",
            "scheduled_trips",
            "avg_scheduled_speed_kmh",
            "median_headway_min",
        ],
        ["Rank", "Route", "Name", "Score", "Trips", "Speed", "Headway"],
        [0.55, 0.6, 1.9, 0.75, 0.7, 0.8, 0.8],
        font_size=8,
    )

    add_paragraph(
        doc,
        "The score bands show that the system has a wide spread of scheduled performance. "
        "In this scoring version, 7 routes fall in the high band, 22 in moderate, 17 in low, and 20 in very low."
    )

    add_section_heading(doc, "Routes for Advocacy Review")
    add_paragraph(
        doc,
        "The advocacy priority score is intentionally different from the efficiency rank. It prioritizes routes with weaker efficiency scores and meaningful service volume, because improvements on those routes may affect more riders."
    )
    add_table_from_df(
        doc,
        priority.head(10),
        [
            "advocacy_priority_rank",
            "route_short_name",
            "route_long_name",
            "advocacy_priority_score",
            "final_efficiency_score",
            "scheduled_trips",
        ],
        ["Priority", "Route", "Name", "Priority score", "Efficiency", "Trips"],
        [0.7, 0.65, 2.2, 1.0, 0.9, 0.75],
        font_size=8,
    )

    add_paragraph(
        doc,
        "The highest-priority review routes are not automatically the worst routes. They are routes where the schedule suggests an advocacy conversation could be productive: "
        "speed may be low, headways may be long, stop density may be high, or directness may be weak while the route still carries substantial scheduled service."
    )

    add_section_heading(doc, "Key Findings")
    add_bullets(
        doc,
        [
            f"Service availability is broad but uneven. The feed contains {summary['total_trips']:,} scheduled weekday trips across 66 active routes, and 52 routes meet the all-day service flag.",
            f"Scheduled speed varies sharply. Route-level average scheduled speed ranges from {speed['avg_scheduled_speed_kmh'].min():.2f} km/h to {speed['avg_scheduled_speed_kmh'].max():.2f} km/h.",
            f"Frequency is not evenly distributed. The median route-level headway is {frequency['median_headway_min'].median():.2f} minutes, but some routes have median headways above 40 minutes.",
            "Several important corridors show up in the advocacy priority lens, including Confederation, Credit Woodlands, Bristol, Hurontario, Cawthra, Bloor, and North Service Road.",
            "School routes and special-purpose routes require separate interpretation because they are not designed to behave like all-day grid routes.",
        ],
    )

    doc.add_page_break()
    add_section_heading(doc, "Advocacy Implications")
    add_paragraph(
        doc,
        "For transit and housing advocacy, the most useful conclusion is not only which route ranks first. The stronger question is where better service would unlock access to housing, campus, jobs, groceries, and late-day travel."
    )
    add_numbered(
        doc,
        [
            "Use the priority list to identify candidate corridors for deeper public-facing analysis, especially routes with high trip volume and weak efficiency indicators.",
            "Ask MiWay and the City of Mississauga for ridership, crowding, on-time performance, pass-up, and stop-level boarding data so scheduled efficiency can be compared with lived rider experience.",
            "Pair route efficiency results with housing density, student housing locations, major employment areas, and late-night travel needs.",
            "Evaluate whether priority routes need bus priority, stop consolidation, branch simplification, schedule changes, or better transfer coordination.",
            "Separate school-special and limited-purpose routes from all-day local routes when communicating results, because their service design goals differ.",
        ],
    )

    add_section_heading(doc, "Limitations")
    add_bullets(
        doc,
        [
            "The analysis uses scheduled GTFS data, not real-time vehicle movement or actual traffic conditions.",
            "The ranking does not include ridership, crowding, fare equity, accessibility barriers, safety, or on-time reliability.",
            "A single representative weekday keeps the project clean, but weekend and seasonal patterns should be studied separately.",
            "The final score is transparent but still depends on chosen weights. The weights should be revisited with stakeholders before using the score for formal policy recommendations.",
            "Lower stop density can improve speed, but stop removal can harm riders with mobility constraints or long walks. Any stop-spacing recommendation needs an equity review.",
        ],
    )

    add_section_heading(doc, "Recommended Next Steps")
    add_numbered(
        doc,
        [
            "Create a map-based appendix showing the top advocacy-priority routes and nearby housing or campus destinations.",
            "Build a sensitivity test that compares alternate scoring weights, such as a rider-access score versus an operating-efficiency score.",
            "Add weekend notebooks for Saturday and Sunday to separate commuter, student, retail, and shift-worker travel patterns.",
            "Compare scheduled speed with real-time travel time data if vehicle location or historical performance feeds become available.",
            "Turn the final route table into a short public brief with 3 to 5 corridor-specific asks.",
        ],
    )

    add_section_heading(doc, "Appendix: Lower-Scoring Routes")
    add_paragraph(
        doc,
        "The routes below have the lowest final scores in this first scoring version. These results should be interpreted with service purpose in mind. A school route, loop route, or coverage route may score low while still serving a valid equity or access function."
    )
    lowest = ranking.sort_values("efficiency_rank", ascending=False).head(10).copy()
    add_table_from_df(
        doc,
        lowest,
        [
            "efficiency_rank",
            "route_short_name",
            "route_long_name",
            "final_efficiency_score",
            "scheduled_trips",
            "avg_scheduled_speed_kmh",
            "median_headway_min",
        ],
        ["Rank", "Route", "Name", "Score", "Trips", "Speed", "Headway"],
        [0.55, 0.6, 1.9, 0.75, 0.7, 0.8, 0.8],
        font_size=8,
    )

    add_section_heading(doc, "Appendix: Files Produced")
    add_bullets(
        doc,
        [
            "Notebooks/01_miway_gtfs_exploration.ipynb",
            "Notebooks/02_miway_route_geometry_metrics.ipynb",
            "Notebooks/03_miway_scheduled_speed_travel_time.ipynb",
            "Notebooks/04_miway_frequency_service_availability.ipynb",
            "Notebooks/05_miway_final_route_efficiency_ranking.ipynb",
            "Outputs/final_route_efficiency_ranking_2026-05-05.csv",
            "Outputs/advocacy_priority_routes_2026-05-05.csv",
            "Outputs/route_efficiency_score_components_2026-05-05.csv",
        ],
    )

    doc.core_properties.title = "MiWay Route Efficiency Analysis"
    doc.core_properties.subject = "GTFS-based route efficiency and advocacy priority report"
    doc.core_properties.author = "Codex for UTM Transit Housing Advocacy"

    doc.save(report_path)
    return report_path


if __name__ == "__main__":
    path = build_report(SAFE_REPORT_PATH)
    print(path)
