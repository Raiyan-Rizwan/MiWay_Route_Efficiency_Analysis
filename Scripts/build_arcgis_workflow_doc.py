from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


PROJECT_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_PATH = PROJECT_ROOT / "Reports" / "ArcGIS_MiWay_Workflow_Pipeline.docx"


BLUE = RGBColor(31, 77, 120)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
BORDER = "D0D7DE"
INK = RGBColor(20, 31, 43)
MUTED = RGBColor(90, 100, 112)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = BORDER) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = margins.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_width(table, widths: list[float]) -> None:
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = Inches(width)
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(sum(widths) * 1440)))
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def style_table(table, widths: list[float], header_rows: int = 1) -> None:
    table.autofit = False
    set_table_width(table, widths)
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cell)
            set_cell_margins(cell)
            if r_idx < header_rows:
                set_cell_shading(cell, LIGHT_BLUE)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = BLUE
            else:
                set_cell_shading(cell, "FFFFFF")


def add_title(document: Document) -> None:
    title = document.add_paragraph()
    title.style = "Title"
    run = title.add_run("MiWay ArcGIS Pro Workflow Pipeline")
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor(11, 37, 69)

    subtitle = document.add_paragraph()
    subtitle.style = "Subtitle"
    run = subtitle.add_run(
        "Creating route lines, joining efficiency results, designing maps, and exporting high-quality layouts"
    )
    run.font.size = Pt(11)
    run.font.color.rgb = MUTED

    meta = document.add_paragraph()
    meta.paragraph_format.space_after = Pt(14)
    run = meta.add_run(
        "Project: UTM Transit Housing Advocacy / MiWay Route Efficiency | Source data: MiWay GTFS and generated route efficiency outputs"
    )
    run.font.size = Pt(9.5)
    run.font.color.rgb = MUTED


def add_callout(document: Document, title: str, body: str) -> None:
    table = document.add_table(rows=1, cols=1)
    style_table(table, [6.5], header_rows=0)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(2)
    title_run = paragraph.add_run(title)
    title_run.bold = True
    title_run.font.color.rgb = BLUE
    body_para = cell.add_paragraph(body)
    body_para.paragraph_format.space_after = Pt(0)
    for run in body_para.runs:
        run.font.size = Pt(9.5)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_heading(document: Document, text: str, level: int = 1) -> None:
    paragraph = document.add_heading(text, level=level)
    for run in paragraph.runs:
        run.font.color.rgb = BLUE if level < 3 else RGBColor(45, 64, 89)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(item)


def add_key_value_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Details"
    for key, value in rows:
        cells = table.add_row().cells
        cells[0].text = key
        cells[1].text = value
    style_table(table, [2.0, 4.5])


def add_workflow_table(document: Document) -> None:
    rows = [
        (
            "1",
            "Create MiWay_Shape_Lines",
            "Use GTFS shape points to create route path geometry.",
            "A line feature class with shapeidtxt, shapeidnum, pointcount, Shape_Length.",
        ),
        (
            "2",
            "Join trips.txt",
            "Attach route_id and direction_id to each shape. Use matching field types.",
            "MiWay_Shape_Lines_With_Routes.",
        ),
        (
            "3",
            "Join routes.txt",
            "Attach route_short_name and route_long_name for readable labels.",
            "MiWay_Route_Lines.",
        ),
        (
            "4",
            "Join efficiency CSV",
            "Attach final_efficiency_score, advocacy_priority_score, ranks, speed, headway, stop spacing, and directness metrics.",
            "MiWay_Route_Efficiency_Lines.",
        ),
        (
            "5",
            "Design map layers",
            "Create one layer for overall efficiency and one layer for advocacy priority.",
            "Route Efficiency Score and Advocacy Priority Score layers.",
        ),
        (
            "6",
            "Build layouts",
            "Create separate landscape layouts for the full network and top 10 advocacy map.",
            "Export-ready map layouts.",
        ),
        (
            "7",
            "Export",
            "Export PDF and PNG with high-resolution settings.",
            "Final report-ready map files.",
        ),
    ]
    table = document.add_table(rows=1, cols=4)
    headers = ["Step", "Task", "Purpose", "Output"]
    for idx, header in enumerate(headers):
        table.cell(0, idx).text = header
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    style_table(table, [0.45, 1.55, 2.4, 2.1])


def add_export_table(document: Document) -> None:
    rows = [
        ("PDF", "300 DPI", "Best", "Embed fonts on; use for submission or portfolio."),
        ("PNG", "300 DPI", "24-bit True Color", "Good for normal report insertion."),
        ("PNG", "600 DPI", "24-bit True Color", "Use when labels or route lines look soft in Word or slides."),
        ("JPEG", "Avoid", "N/A", "JPEG compression makes labels and thin linework blurry."),
    ]
    table = document.add_table(rows=1, cols=4)
    headers = ["Format", "Resolution", "Quality / Color", "Use"]
    for idx, header in enumerate(headers):
        table.cell(0, idx).text = header
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    style_table(table, [1.0, 1.1, 1.65, 2.75])


def configure_styles(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for style_name, size, before, after in [
        ("Heading 1", 16, 14, 6),
        ("Heading 2", 13, 10, 4),
        ("Heading 3", 11.5, 8, 3),
    ]:
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLUE
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.15


def build_document() -> None:
    document = Document()
    configure_styles(document)
    add_title(document)

    add_callout(
        document,
        "Core idea",
        "GTFS splits the map path, route identity, route names, and analysis scores into different files. The ArcGIS workflow reconnects those pieces into one final map-ready line layer.",
    )

    add_heading(document, "Pipeline Overview")
    add_workflow_table(document)

    add_heading(document, "Input Files")
    add_key_value_table(
        document,
        [
            ("MiWay_Shape_Lines", "Line geometry created from GTFS shapes. Initial fields include ObjectID, Shape, Shape_Length, shapeidtxt, shapeidnum, and pointcount."),
            ("trips.txt", "Connects each shape_id to route_id and direction_id. This tells ArcGIS which route uses each shape."),
            ("routes.txt", "Adds readable route information such as route_short_name and route_long_name."),
            ("final_route_efficiency_ranking_2026-05-05.csv", "Adds final_efficiency_score, advocacy_priority_score, advocacy_priority_rank, scheduled speed, headway, directness, and stop spacing metrics."),
            ("arcgis_stops_clean.csv", "Stop point table used to create MiWay stop points with XY Table To Point."),
        ],
    )

    add_heading(document, "Why The Joins Were Necessary")
    paragraph = document.add_paragraph()
    paragraph.add_run("MiWay_Shape_Lines only contains geometry. ").bold = True
    paragraph.add_run(
        "It knows where each shape goes, but it does not know the bus route name, route score, or advocacy priority score. The joins add that missing context in stages."
    )
    add_bullets(
        document,
        [
            "shapes.txt / MiWay_Shape_Lines = where the bus route goes.",
            "trips.txt = which route uses that shape.",
            "routes.txt = what the route is called.",
            "efficiency CSV = how well that route performs and how urgent it is for advocacy.",
        ],
    )

    add_heading(document, "Step 1: Create MiWay_Shape_Lines")
    document.add_paragraph(
        "Create lines from GTFS shape points by grouping points by shape_id and ordering them by shape_pt_sequence. The resulting layer should contain line geometry and fields such as shapeidtxt, shapeidnum, pointcount, and Shape_Length."
    )
    add_callout(
        document,
        "Quality check",
        "Open the attribute table. If you see one row per shape line and the map draws route paths across Mississauga, the geometry step worked.",
    )

    add_heading(document, "Step 2: Join trips.txt To Add route_id")
    add_numbered(
        document,
        [
            "Right-click MiWay_Shape_Lines.",
            "Choose Joins and Relates, then Add Join.",
            "Set Input Table to MiWay_Shape_Lines.",
            "Set Input Join Field to shapeidnum if trips.shape_id is numeric. Use shapeidtxt only if trips.shape_id is text.",
            "Set Join Table to trips.txt.",
            "Set Join Table Field to shape_id.",
            "Open the attribute table and scroll to the far right to confirm route_id and direction_id appeared.",
        ],
    )
    add_callout(
        document,
        "If ArcGIS says value type is incompatible",
        "The two join fields are different data types. Use shapeidnum when trips.shape_id is numeric. Use shapeidtxt when trips.shape_id is text.",
    )

    add_heading(document, "Step 3: Export The Joined Shape Layer")
    document.add_paragraph(
        "ArcGIS joins are temporary. After confirming the trips.txt join worked, export the layer so route_id and direction_id become permanent fields."
    )
    add_bullets(
        document,
        [
            "Right-click MiWay_Shape_Lines.",
            "Choose Data, then Export Features.",
            "Save the output as MiWay_Shape_Lines_With_Routes.",
        ],
    )

    add_heading(document, "Step 4: Join routes.txt To Add Route Names")
    add_numbered(
        document,
        [
            "Right-click MiWay_Shape_Lines_With_Routes.",
            "Choose Joins and Relates, then Add Join.",
            "Use route_id as the input join field.",
            "Use routes.txt as the join table.",
            "Use route_id as the join table field.",
            "Confirm fields such as route_short_name and route_long_name appear in the attribute table.",
            "Export Features and save as MiWay_Route_Lines.",
        ],
    )

    add_heading(document, "Step 5: Join The Route Efficiency CSV")
    add_numbered(
        document,
        [
            "Right-click MiWay_Route_Lines.",
            "Choose Joins and Relates, then Add Join.",
            "Use route_id as the input join field.",
            "Use final_route_efficiency_ranking_2026-05-05.csv as the join table.",
            "Use route_id as the join table field.",
            "Confirm fields such as final_efficiency_score, advocacy_priority_score, advocacy_priority_rank, avg_scheduled_speed_kmh, median_headway_min, avg_directness_ratio, and stops_per_km appear.",
            "Export Features and save as MiWay_Route_Efficiency_Lines.",
        ],
    )

    add_heading(document, "Step 6: Symbolize Route Efficiency")
    add_key_value_table(
        document,
        [
            ("Layer", "MiWay_Route_Efficiency_Lines, renamed to Route Efficiency Score."),
            ("Primary symbology", "Graduated Colors."),
            ("Field", "final_efficiency_score."),
            ("Normalization", "None. The score is already a comparable index."),
            ("Method", "Natural Breaks (Jenks)."),
            ("Classes", "5."),
            ("Color scheme", "Red-Yellow-Green, where low scores are red/orange and high scores are green."),
        ],
    )

    add_heading(document, "Step 7: Create The Advocacy Priority Layer")
    add_numbered(
        document,
        [
            "In the Contents pane, right-click Route Efficiency Score and copy it.",
            "Right-click the map name and paste.",
            "Rename the copied layer Advocacy Priority Score.",
            "Open Symbology and keep Graduated Colors.",
            "Set Field to advocacy_priority_score.",
            "Set Normalization to None.",
            "Set Method to Natural Breaks (Jenks).",
            "Set Classes to 5 for the full network map or 3 for a cleaner top 10 map.",
            "Use a Yellow-Orange-Red color scheme, where high priority appears red.",
        ],
    )

    add_heading(document, "Step 8: Filter Top 10 Advocacy Priority Routes")
    document.add_paragraph(
        "Apply the top 10 filter only to the Advocacy Priority Score layer, not to the full efficiency layer."
    )
    add_bullets(
        document,
        [
            "Right-click Advocacy Priority Score.",
            "Choose Properties, then Definition Query.",
            "Create a query using advocacy_priority_rank <= 10.",
            "If the field name became long after joining, use the field picker and choose the field ending in advocacy_priority_rank.",
        ],
    )

    add_heading(document, "Step 9: Add Stops")
    add_numbered(
        document,
        [
            "Add Outputs\\arcgis_stops_clean.csv to the map.",
            "Open Geoprocessing and search for XY Table To Point.",
            "Set Input Table to arcgis_stops_clean.csv.",
            "Set X Field to stop_lon.",
            "Set Y Field to stop_lat.",
            "Set Coordinate System to WGS 1984.",
            "Save the output feature class as MiWay_Stops.",
            "Symbolize stops as small gray or dark dots so they support the line map without overpowering it.",
        ],
    )

    add_heading(document, "Step 10: Build Two Separate Maps")
    document.add_paragraph(
        "The safest approach is to duplicate the map before creating layouts. This prevents one layout from changing when layer visibility is adjusted for another layout."
    )
    add_bullets(
        document,
        [
            "In the Catalog pane, expand Maps.",
            "Copy the current map and paste it twice if needed.",
            "Rename one map MiWay Efficiency Map.",
            "Rename the other map MiWay Advocacy Priority Map.",
            "In MiWay Efficiency Map, show Route Efficiency Score, MiWay_Stops, and the basemap.",
            "In MiWay Advocacy Priority Map, show Advocacy Priority Score with the top 10 definition query, MiWay_Stops, and the basemap.",
        ],
    )

    add_heading(document, "Step 11: Fit Mississauga In The Layout")
    add_numbered(
        document,
        [
            "Go to Insert, then New Layout, and choose Letter Landscape.",
            "Insert a Map Frame and choose the relevant map.",
            "Resize the map frame so it fills most of the page while leaving room for the title and legend.",
            "Right-click the map frame and choose Activate.",
            "Right-click the route layer and choose Zoom To Layer.",
            "Pan and zoom until Mississauga fills the map frame with a small margin around the routes.",
            "Click Layout, then Close Activation.",
        ],
    )
    add_callout(
        document,
        "Important layout concept",
        "Resizing the map frame changes the size of the picture box. Activating the map frame changes the geographic area shown inside that box. To fit Mississauga, activate the map frame and adjust the extent.",
    )

    add_heading(document, "Step 12: Add Layout Elements")
    add_bullets(
        document,
        [
            "Add a title. Suggested title for Map 1: MiWay Route Efficiency Score.",
            "Add a title. Suggested title for Map 2: Top 10 MiWay Advocacy Priority Routes.",
            "Add a legend and remove unnecessary layers from it.",
            "Add a north arrow and scale bar.",
            "Add a data source note: Source: MiWay GTFS schedule data and route efficiency analysis, May 2026.",
        ],
    )

    add_heading(document, "Step 13: Export High-Quality Layouts")
    document.add_paragraph(
        "Use Share, then Export Layout. Do not screenshot the layout. PDF usually preserves labels and linework better because much of the map remains vector-based."
    )
    add_export_table(document)

    add_heading(document, "Suggested Output Names")
    add_bullets(
        document,
        [
            "Outputs\\ArcGIS\\miway_route_efficiency_score_map.pdf",
            "Outputs\\ArcGIS\\miway_route_efficiency_score_map.png",
            "Outputs\\ArcGIS\\miway_top_10_advocacy_priority_routes.pdf",
            "Outputs\\ArcGIS\\miway_top_10_advocacy_priority_routes.png",
        ],
    )

    add_heading(document, "Troubleshooting Notes")
    add_key_value_table(
        document,
        [
            ("Join appears missing", "The layer name does not change after Add Join. Reopen the attribute table and scroll all the way right to see joined fields."),
            ("No new layer appears", "Add Join is temporary and modifies the existing layer view. Export Features creates the permanent layer."),
            ("Incompatible field types", "Use shapeidnum for numeric shape_id joins and shapeidtxt for text shape_id joins."),
            ("Field names are long", "After joins, ArcGIS may prefix field names with table names. Use the field picker when writing definition queries."),
            ("Map frame does not fit Mississauga", "Activate the map frame and zoom to the route layer. Do not fix extent by only resizing the frame."),
            ("Export looks blurry", "Export Layout as PDF at 300 DPI with best image quality, or PNG at 600 DPI. Avoid JPEG."),
        ],
    )

    add_heading(document, "Final Checklist")
    add_bullets(
        document,
        [
            "MiWay_Route_Efficiency_Lines contains route_id, route name fields, final_efficiency_score, advocacy_priority_score, and advocacy_priority_rank.",
            "Route Efficiency Score uses final_efficiency_score with Natural Breaks, 5 classes, no normalization, and a red-yellow-green scheme.",
            "Advocacy Priority Score uses advocacy_priority_score with no normalization and a yellow-orange-red scheme.",
            "Top 10 advocacy map has advocacy_priority_rank <= 10 applied.",
            "Stops are created from stop_lon and stop_lat using WGS 1984.",
            "Each layout has a clean title, legend, north arrow, scale bar, and data source note.",
            "Final exports include both PDF and PNG versions.",
        ],
    )

    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run("MiWay Route Efficiency | ArcGIS Pro Workflow")
    footer_run.font.size = Pt(8.5)
    footer_run.font.color.rgb = MUTED

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT_PATH)


if __name__ == "__main__":
    build_document()
    print(OUTPUT_PATH)
