from copy import deepcopy
from pathlib import Path

from docx import Document


SOURCE = Path(r"C:\Users\Raiyan Rizwan\Desktop\Job Application Automation Platform\Raiyan_Bin_Rizwan_Resume_Updated.docx")
OUTPUT = SOURCE.with_name("Raiyan_Bin_Rizwan_Resume_Updated2.docx")


def copy_paragraph_format(src, dst):
    dst.style = src.style
    if src._p.pPr is not None:
        if dst._p.pPr is not None:
            dst._p.remove(dst._p.pPr)
        dst._p.insert(0, deepcopy(src._p.pPr))


def clear_runs(paragraph):
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)


def replace_plain(paragraph, text):
    template_run = paragraph.runs[0] if paragraph.runs else None
    clear_runs(paragraph)
    run = paragraph.add_run(text)
    if template_run is not None:
        run.bold = template_run.bold
        run.italic = template_run.italic
        run.underline = template_run.underline
        run.font.name = template_run.font.name
        run.font.size = template_run.font.size


def replace_heading_with_date(paragraph, title, date_text):
    template_title = paragraph.runs[0] if paragraph.runs else None
    template_date = paragraph.runs[1] if len(paragraph.runs) > 1 else None
    clear_runs(paragraph)
    title_run = paragraph.add_run(title)
    if template_title is not None:
        title_run.bold = template_title.bold
        title_run.italic = template_title.italic
        title_run.font.name = template_title.font.name
        title_run.font.size = template_title.font.size
    date_run = paragraph.add_run("\t" + date_text)
    if template_date is not None:
        date_run.bold = template_date.bold
        date_run.italic = template_date.italic
        date_run.font.name = template_date.font.name
        date_run.font.size = template_date.font.size


def make_like_before(anchor, template, text):
    new_p = anchor.insert_paragraph_before()
    copy_paragraph_format(template, new_p)
    replace_plain(new_p, text)
    return new_p


def make_heading_like_before(anchor, template, title, date_text):
    new_p = anchor.insert_paragraph_before()
    copy_paragraph_format(template, new_p)
    replace_heading_with_date(new_p, title, date_text)
    return new_p


doc = Document(SOURCE)
paragraphs = doc.paragraphs

replacements = {
    "Architected an end-to-end Azure Custom Speech training pipeline in Node.js, enabling automatic acoustic model fine-tuning on clinical audio data — one of few student-built systems integrating real clinical speech AI at this depth":
        "Delivered a student-built clinical speech AI retraining system that automated acoustic model fine-tuning for personalized medical transcription workflows.",
    "Visualized protein structures and drug binding sites using Python and Amazon AWS cloud computing services":
        "Enabled bioinformatics researchers to explore protein structures and drug-binding patterns through cloud-hosted Python/SQL visualization tools.",
    "Provide exceptional customer service by assisting guests, adults and children, with arcade game instructions and ensuring a positive gaming experience across the facility":
        "Supported a safe, positive guest experience across high-traffic arcade and activity areas by combining customer service, safety communication, and operations support.",
    "Built an end-to-end automated trading pipeline on all 500+ S&P 500 constituents (2018–2026), engineering 6 technical features per stock (Garman-Klass Volatility, RSI, Bollinger Bands, ATR, MACD, Dollar Volume) from raw OHLCV data via yfinance.":
        "Created a research-grade strategy pipeline to test whether unsupervised learning could identify tradable S&P 500 stock regimes across 2018–2026.",
    "Analysed 19,364 respondents across 10 countries from WVS Wave 7 using PCA, factor analysis (varimax), GMM clustering, and LDA to identify three latent value dimensions: traditional engagement, subjective well-being, and social openness.":
        "Produced a cross-country statistical analysis showing that 19,364 World Values Survey respondents clustered around three interpretable latent value dimensions.",
    "Built a cloud-hosted platform for visualizing protein structures and drug binding sites, used in bioinformatics research workflows.":
        "Built a data analysis workflow to examine how COVID-19 affected Canadian businesses, translating business-impact data into Python/SQL-backed visualizations.",
    "Designed SQL queries and Python scripts to power interactive HTML dashboards, enabling non-technical researchers to explore structural data.":
        "Designed SQL queries, Python scripts, and HTML/CSS dashboards to make business-impact trends easier to explore and communicate.",
}

for paragraph in paragraphs:
    text = paragraph.text
    if text in replacements:
        replace_plain(paragraph, replacements[text])
    elif text.startswith("Affect of Covid-19 on Canadian Businesses"):
        replace_heading_with_date(
            paragraph,
            "Effect of COVID-19 on Canadian Businesses (Supervised by STA305 Prof.)",
            "Jun 2024 – Aug 2025",
        )
    elif text == "Tools: Git, ffmpeg, Express.js, HTML/CSS":
        # Preserve bold label + regular value formatting.
        clear_runs(paragraph)
        label = paragraph.add_run("Tools: ")
        label.bold = True
        label.font.name = "Arial"
        value = paragraph.add_run("Git, ArcGIS Pro, spatial analysis, cartographic design, ffmpeg, Express.js, HTML/CSS")
        value.font.name = "Arial"

paragraphs = doc.paragraphs
tech_anchor = next(p for p in paragraphs if p.text == "TECHNICAL SKILLS")
project_title_index = next(
    i for i, p in enumerate(paragraphs) if p.text.startswith("Algorithmic Trading Strategy")
)
project_title_template = paragraphs[project_title_index]
project_tools_template = paragraphs[project_title_index + 1]
project_bullet_template = paragraphs[project_title_index + 2]

make_heading_like_before(
    tech_anchor,
    project_title_template,
    "MiWay Route Efficiency & Advocacy Priority Mapping",
    "May 2026",
)
make_like_before(
    tech_anchor,
    project_tools_template,
    "ArcGIS Pro · Spatial Analysis · Transit Planning · Data Visualization · Cartographic Design",
)
miway_bullets = [
    "Mapped MiWay route efficiency to identify underperforming corridors and convert service-performance patterns into evidence-based transit advocacy priorities.",
    "Built route-level efficiency and advocacy priority scores using normalized fields, classification breaks, and graduated line symbology across the Mississauga network.",
    "Produced two ArcGIS map layouts with basemaps, legends, scale bars, and north arrows to communicate citywide route performance and top-priority routes clearly.",
]
for bullet in miway_bullets:
    make_like_before(tech_anchor, project_bullet_template, bullet)

doc.save(OUTPUT)
print(OUTPUT)
